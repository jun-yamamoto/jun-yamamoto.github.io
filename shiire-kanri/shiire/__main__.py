"""コマンドライン: python -m shiire <command> ...

  init                         台帳を新規作成（既存なら何もしない）
  add   --date --name --price  ヤフオク/EC 1件登録（URLは --url）
  text  "キタムラ Nikon F100 仕入38500 送料1000 7/11"   1行入力で実店舗登録
  dashboard                    ダッシュボードの ENTRIES を再生成
  summary [--month N]          月の仕入合計（件数・金額・送料・合計）
  mf    journal.csv [--ids]    MF仕訳帳照合（--ids で B/C の取引番号だけ）
  backup                       台帳をバックアップ
  shot  明細URL --id 商品ID --date 落札日   支払い明細のスクショ（Windows）
"""
from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from . import config, dashboard, freetext, mf_match
from .files import LedgerLockedError, backup
from .ledger import Entry, add_entry, create_ledger, edit_ledger, month_summary


def _date(s: str) -> date:
    parts = [int(p) for p in s.replace("-", "/").split("/")]
    return date(*parts) if len(parts) == 3 else date(config.LEDGER_YEAR, *parts)


def _print_result(res) -> None:
    label = {"added": "登録", "filled": "空欄補完", "duplicate": "登録済み（変更なし）"}[res.status]
    print(f"{label}: {res.sheet} {res.row}行目" + (f"  T番号 {res.invoice}" if res.invoice else ""))
    for m in res.messages:
        print("  - " + m)


def _print_summary(s) -> None:
    print(f"{s.month}月 仕入合計: {s.count}件 / 金額 {s.price:,}円 / 送料 {s.shipping:,}円 / 合計 {s.total:,}円"
          + (f"（返品 {s.returned_total:,}円 除く）" if s.returned_total else ""))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="shiire")
    ap.add_argument("--base", type=Path, default=config.BASE_DIR, help="作業フォルダ")
    sub = ap.add_subparsers(dest="cmd", required=True)

    sub.add_parser("init")
    sub.add_parser("backup")
    sub.add_parser("dashboard")
    p = sub.add_parser("summary")
    p.add_argument("--month", type=int, default=date.today().month)

    p = sub.add_parser("add")
    p.add_argument("--date", type=_date, required=True, help="落札日（終了日） YYYY/MM/DD")
    p.add_argument("--name", required=True)
    p.add_argument("--price", type=int, required=True, help="落札金額（クーポン前）")
    p.add_argument("--coupon", type=int, default=0)
    p.add_argument("--shipping", type=int)
    p.add_argument("--shop", default="")
    p.add_argument("--address", default="")
    p.add_argument("--payment", default="", help="AMEX / PAYPAY（cards.json で展開）または実値")
    p.add_argument("--type", dest="shop_type", choices=config.SHOP_TYPES, default="ヤフオク")
    p.add_argument("--invoice", default="")
    p.add_argument("--url", dest="evidence", default="", help="商品URL（証憑）")
    p.add_argument("--note", default="")
    for flag in ("returned", "repair", "junk"):
        p.add_argument(f"--{flag}", action="store_true")

    p = sub.add_parser("text")
    p.add_argument("line")

    p = sub.add_parser("mf")
    p.add_argument("csv", type=Path)
    p.add_argument("--ids", action="store_true", help="B/C の取引番号だけをドット区切りで出力")
    p.add_argument("--days", type=int, default=7, help="日付差の許容日数")

    # Windows のパイプ出力（cp932）で文字化け・エラーにならないよう UTF-8 に固定
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    p = sub.add_parser("shot")
    p.add_argument("url", help="かんたん決済 支払い明細のURL")
    p.add_argument("--id", required=True, help="商品ID（例 x1234567890）")
    p.add_argument("--date", type=_date, required=True, help="落札日")

    a = ap.parse_args(argv)
    path = config.ledger_path(a.base)

    if a.cmd == "init":
        if path.exists():
            print(f"既に存在します: {path}")
        else:
            create_ledger(path)
            print(f"作成しました: {path}")
    elif a.cmd == "backup":
        print(backup(path) or "台帳がありません")
    elif a.cmd in ("add", "text"):
        if a.cmd == "text":
            entry = freetext.parse_line(a.line, base=a.base)
        else:
            fields = {k: getattr(a, k) for k in Entry.__dataclass_fields__ if hasattr(a, k)}
            fields["date"], fields["name"] = a.date, a.name
            fields["payment"] = config.resolve_payment(a.payment, a.base)
            entry = Entry(**fields)
        with edit_ledger(path) as wb:
            res = add_entry(wb, entry)
        _print_result(res)
        _print_summary(month_summary(load_workbook(path), entry.date.month))
    elif a.cmd == "dashboard":
        n = dashboard.regenerate(load_workbook(path), a.base / config.DASHBOARD_NAME)
        print(f"ダッシュボード更新: {n}件")
    elif a.cmd == "summary":
        _print_summary(month_summary(load_workbook(path), a.month))
    elif a.cmd == "mf":
        wb = load_workbook(path)
        ms = mf_match.match(wb, mf_match.load_journals_csv(a.csv, config.LEDGER_YEAR), day_tolerance=a.days)
        print(mf_match.numbers_only(ms) if a.ids else mf_match.report(ms))
    elif a.cmd == "shot":
        from .screenshot_win import capture_payment_detail

        print(f"保存: {capture_payment_detail(a.url, a.base, a.id, a.date)}")
        print("→ PNG を開いて 商品ID・支払金額合計・支払方法 を目視確認してください")
    return 0


def run() -> int:
    try:
        return main()
    except (KeyError, ValueError, FileNotFoundError, LedgerLockedError) as e:
        print(f"エラー: {e.args[0] if e.args else e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(run())
