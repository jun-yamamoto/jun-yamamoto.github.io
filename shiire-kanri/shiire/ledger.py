"""仕入台帳（openpyxl）の作成・行追加・検索。

破損防止のため、次の点を必ず守る:
- 行挿入前に合計行・返品行の A:B 結合を unmerge → 挿入後に再 merge
- insert_rows は数式を調整しないので SUMIF 範囲を書き直す
- 縞模様は挿入後に行番号の偶奇で塗り直す
- データ検証・条件付き書式・タブ色は ensure_sheet_rules で再設定する
"""
from __future__ import annotations

import re
import unicodedata
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterator

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from . import config
from .files import backup, check_locks

C = config.COL
NCOLS = len(config.HEADERS)
LAST_COL = get_column_letter(NCOLS)  # O
MONEY_FMT = "#,##0"
DATE_FMT = "yyyy/mm/dd"
ITEM_ID_RE = re.compile(r"/auction/([A-Za-z]?\d{6,})")


# ---------------------------------------------------------------- データ型

@dataclass
class Entry:
    """台帳1行分。price はクーポン適用前の落札額/仕入額、coupon は値引額。"""

    date: date
    name: str
    price: int
    shipping: int | None = None
    shop: str = ""
    address: str = ""
    payment: str = ""
    shop_type: str = "ヤフオク"
    invoice: str = ""
    evidence: str = ""
    note: str = ""
    coupon: int = 0
    returned: bool = False
    repair: bool = False
    junk: bool = False

    @property
    def net_price(self) -> int:
        """金額列に書く値（クーポン後）。"""
        return self.price - (self.coupon or 0)

    @property
    def item_id(self) -> str | None:
        return extract_item_id(self.evidence)


@dataclass
class Row:
    sheet: str
    row: int
    values: dict = field(default_factory=dict)

    def __getitem__(self, key):
        return self.values.get(key)


@dataclass
class AddResult:
    status: str  # "added" | "filled" | "duplicate"
    sheet: str
    row: int
    invoice: str = ""
    messages: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- 文字列処理

def extract_item_id(text: str | None) -> str | None:
    if not text:
        return None
    m = ITEM_ID_RE.search(str(text))
    return m.group(1) if m else None


def normalize_shop(name: str | None) -> str:
    """表記ゆれ吸収用の正規化（全半角・空白・法人格・店舗接尾辞）。"""
    s = unicodedata.normalize("NFKC", name or "").lower()
    for w in ("株式会社", "(株)", "有限会社", "(有)", "合同会社"):
        s = s.replace(w, "")
    s = re.sub(r"[\s・\-‐ー_/　]+", "", s)
    return s


def same_shop(a: str | None, b: str | None) -> bool:
    na, nb = normalize_shop(a), normalize_shop(b)
    if not na or not nb:
        return False
    if na == nb:
        return True
    shorter, longer = sorted((na, nb), key=len)
    return len(shorter) >= 3 and shorter in longer


def sanitize_invoice(invoice: str | None) -> str:
    s = unicodedata.normalize("NFKC", invoice or "").strip().upper().replace("-", "")
    if not s:
        return ""
    if not re.fullmatch(r"T\d{13}", s):
        return s  # 形式外は触らない（人が確認する）
    return "" if s in config.FORBIDDEN_INVOICES else s


def fixed_invoice(shop: str) -> str | None:
    n = normalize_shop(shop)
    for keywords, t in config.FIXED_INVOICES:
        if any(normalize_shop(k) in n for k in keywords):
            return t
    return None


# ---------------------------------------------------------------- 書式

def _font(bold=False, color=None) -> Font:
    return Font(name=config.FONT_NAME, size=config.FONT_SIZE, bold=bold, color=color)


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def style_data_row(ws: Worksheet, r: int) -> None:
    fill = _fill(config.STRIPE_ODD if r % 2 else config.STRIPE_EVEN)
    for c in range(1, NCOLS + 1):
        cell = ws.cell(r, c)
        cell.fill = fill
        cell.font = _font()
        cell.alignment = Alignment(vertical="center",
                                   horizontal="center" if c >= C["返品"] else None)
    ws.cell(r, C["日付"]).number_format = DATE_FMT
    for key in ("金額", "送料", "合計金額"):
        ws.cell(r, C[key]).number_format = MONEY_FMT


def style_total_row(ws: Worksheet, r: int) -> None:
    for c in range(1, NCOLS + 1):
        cell = ws.cell(r, c)
        cell.fill = _fill(config.TOTAL_FILL)
        cell.font = _font(bold=True, color="FFFFFFFF")
        if c in (C["金額"], C["送料"], C["合計金額"]):
            cell.number_format = MONEY_FMT
    ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")


def ensure_sheet_rules(ws: Worksheet, last_row: int | None = None) -> None:
    """✓ドロップダウン・条件付き書式・タブ色を（再）設定する。copy_worksheet 後にも呼ぶ。"""
    end = max(config.RULE_LAST_ROW, (last_row or 0) + 50)
    ws.data_validations.dataValidation = []
    dv = DataValidation(type="list", formula1='"✓,"', allow_blank=True)
    dv.add(f"M{config.FIRST_DATA_ROW}:O{end}")
    ws.add_data_validation(dv)

    ws.conditional_formatting = ConditionalFormattingList()
    rng = f"A{config.FIRST_DATA_ROW}:{LAST_COL}{end}"
    r0 = config.FIRST_DATA_ROW
    ret = FormulaRule(formula=[f'$M{r0}="{config.CHECK}"'], stopIfTrue=True,
                      fill=PatternFill(fill_type="solid", bgColor=config.RETURN_BG),
                      font=Font(color=config.RETURN_FG))
    rep = FormulaRule(formula=[f'$N{r0}="{config.CHECK}"'], stopIfTrue=True,
                      fill=PatternFill(fill_type="solid", bgColor=config.REPAIR_BG),
                      font=Font(color=config.REPAIR_FG))
    ws.conditional_formatting.add(rng, ret)  # priority 1
    ws.conditional_formatting.add(rng, rep)  # priority 2
    ret.priority, rep.priority = 1, 2
    ws.sheet_properties.tabColor = config.TAB_MONTH


# ---------------------------------------------------------------- 作成

def _write_totals(ws: Worksheet, total_row: int) -> None:
    """合計行・返品行の SUMIF を data 範囲 3..total_row-1 で書き直す。"""
    ret_row = total_row + 1
    n = total_row - 1
    ws.cell(total_row, 1, config.TOTAL_LABEL)
    ws.cell(ret_row, 1, config.RETURN_LABEL)
    for key in ("金額", "送料", "合計金額"):
        col = get_column_letter(C[key])
        ws.cell(total_row, C[key], f'=SUMIF(M3:M{n},"<>{config.CHECK}",{col}3:{col}{n})')
        ws.cell(ret_row, C[key], f'=SUMIF(M3:M{n},"{config.CHECK}",{col}3:{col}{n})')
    for r in (total_row, ret_row):
        ws.merge_cells(f"A{r}:B{r}")
        style_total_row(ws, r)


def _init_month_sheet(ws: Worksheet, year: int, month: int) -> None:
    ws.cell(1, 1, f"{year}年{month}月 仕入台帳")
    ws.merge_cells(f"A1:{LAST_COL}1")
    ws.cell(1, 1).font = Font(name=config.FONT_NAME, size=14, bold=True, color="FF4A148C")
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    for i, h in enumerate(config.HEADERS, start=1):
        cell = ws.cell(config.HEADER_ROW, i, h)
        cell.fill = _fill(config.HEADER_FILL)
        cell.font = _font(bold=True, color="FFFFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = config.COL_WIDTHS[i - 1]
    ws.freeze_panes = "C3"
    # 空の月は空行1行（3行目）＋合計行4・返品行5 で始める（SUMIF範囲が合計行を含まないように）
    r = config.FIRST_DATA_ROW
    ws.cell(r, C["合計金額"], _total_formula(r))
    style_data_row(ws, r)
    _write_totals(ws, r + 1)
    ensure_sheet_rules(ws)


def _total_formula(r: int) -> str:
    return f'=IFERROR(C{r}+IF(D{r}="",0,D{r}),"")'


def _init_annual_sheet(ws: Worksheet, year: int) -> None:
    ws.cell(1, 1, f"{year}年 仕入 年間合計")
    ws.merge_cells("A1:E1")
    ws.cell(1, 1).font = Font(name=config.FONT_NAME, size=14, bold=True, color="FF4A148C")
    heads = ["月", "金額", "送料", "合計金額", "うち返品（合計金額）"]
    for i, h in enumerate(heads, start=1):
        cell = ws.cell(2, i, h)
        cell.fill = _fill(config.HEADER_FILL)
        cell.font = _font(bold=True, color="FFFFFFFF")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = [8, 14, 12, 14, 20][i - 1]
    for m in range(1, 13):
        r = m + 2
        sh = f"'{m}月'"
        ws.cell(r, 1, f"{m}月")
        for i, col in ((2, "C"), (3, "D"), (4, "E")):
            ws.cell(r, i, f'=IFERROR(INDEX({sh}!{col}:{col},MATCH("{config.TOTAL_LABEL}",{sh}!A:A,0)),"")')
        ws.cell(r, 5, f'=IFERROR(INDEX({sh}!E:E,MATCH("{config.RETURN_LABEL}",{sh}!A:A,0)),"")')
        style_data_row(ws, r)
        for c in range(2, 6):
            ws.cell(r, c).number_format = MONEY_FMT
    tr = 15
    ws.cell(tr, 1, config.TOTAL_LABEL)
    for i in range(2, 6):
        col = get_column_letter(i)
        ws.cell(tr, i, f"=SUM({col}3:{col}14)")
    for c in range(1, 6):
        cell = ws.cell(tr, c)
        cell.fill = _fill(config.TOTAL_FILL)
        cell.font = _font(bold=True, color="FFFFFFFF")
        if c > 1:
            cell.number_format = MONEY_FMT
    ws.sheet_properties.tabColor = config.TAB_ANNUAL


def create_ledger(path: Path, year: int = config.LEDGER_YEAR) -> Workbook:
    wb = Workbook()
    annual = wb.active
    annual.title = config.ANNUAL_SHEET
    _init_annual_sheet(annual, year)
    for m in range(1, 13):
        _init_month_sheet(wb.create_sheet(f"{m}月"), year, m)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return wb


@contextmanager
def edit_ledger(path: Path, year: int = config.LEDGER_YEAR, do_backup: bool = True) -> Iterator[Workbook]:
    """ロック確認 → バックアップ → 編集 → 保存。台帳が無ければ新規作成。"""
    check_locks(path)
    if not path.exists():
        create_ledger(path, year)
    elif do_backup:
        backup(path)
    wb = load_workbook(path)
    yield wb
    wb.save(path)


# ---------------------------------------------------------------- 読み取り

def find_label_row(ws: Worksheet, label: str) -> int | None:
    for r in range(config.FIRST_DATA_ROW, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip() == label.strip():
            return r
    return None


def _is_blank_row(ws: Worksheet, r: int) -> bool:
    return all(ws.cell(r, C[k]).value in (None, "") for k in ("日付", "商品名", "金額"))


def iter_rows(wb: Workbook) -> Iterator[Row]:
    """全月シートのデータ行（合計行・空行はスキップ）。"""
    for name in config.MONTH_SHEETS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        end = find_label_row(ws, config.TOTAL_LABEL) or ws.max_row + 1
        for r in range(config.FIRST_DATA_ROW, end):
            if _is_blank_row(ws, r):
                continue
            vals = {h: ws.cell(r, i).value for i, h in enumerate(config.HEADERS, start=1)}
            yield Row(name, r, vals)


def to_date(v) -> date | None:
    """datetime（3月〜）と文字列日付（1-2月）の両方を受け付ける。"""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        m = re.match(r"\s*(\d{4})[/\-.年](\d{1,2})[/\-.月](\d{1,2})", v)
        if m:
            return date(int(m[1]), int(m[2]), int(m[3]))
    return None


def to_int(v) -> int:
    if v in (None, ""):
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    s = re.sub(r"[^\d\-]", "", str(v))
    return int(s) if s not in ("", "-") else 0


def find_by_item_id(wb: Workbook, item_id: str) -> Row | None:
    for row in iter_rows(wb):
        if extract_item_id(row["証憑"]) == item_id:
            return row
    return None


def lookup_shop(wb: Workbook, shop: str) -> tuple[str, str]:
    """当年全シートから同店舗の (T番号, 住所) を流用。新しい行を優先。"""
    invoice = address = ""
    for row in iter_rows(wb):
        if not same_shop(row["店舗名"], shop):
            continue
        t = sanitize_invoice(row["インボイス番号"])
        if t:
            invoice = t
        if row["住所"]:
            address = str(row["住所"])
    return invoice, address


def resolve_invoice(wb: Workbook, shop: str, given: str = "") -> str:
    """固定マッピング > 指定値 > 過去登録流用。LINEヤフーのT番号は除外。"""
    fixed = fixed_invoice(shop)
    if fixed:
        return fixed
    t = sanitize_invoice(given)
    if t:
        return t
    return lookup_shop(wb, shop)[0]


# ---------------------------------------------------------------- 書き込み

def _write_entry(ws: Worksheet, r: int, e: Entry, invoice: str, address: str) -> None:
    note = e.note or ""
    if e.coupon:
        tag = f"クーポン使用 -{e.coupon:,}円"
        if tag not in note:
            note = f"{tag} {note}".strip()
    values = {
        "日付": datetime(e.date.year, e.date.month, e.date.day),
        "商品名": e.name,
        "金額": e.net_price,
        "送料": e.shipping if e.shipping is not None else None,
        "合計金額": _total_formula(r),
        "店舗名": e.shop,
        "住所": address,
        "支払方法": e.payment,
        "店舗タイプ": e.shop_type,
        "インボイス番号": invoice or None,
        "証憑": e.evidence,
        "備考": note or None,
        "返品": config.CHECK if e.returned else None,
        "修理": config.CHECK if e.repair else None,
        "ジャンク": config.CHECK if e.junk else None,
    }
    for k, v in values.items():
        ws.cell(r, C[k], v)


def _append_row(ws: Worksheet) -> int:
    """合計行の直前に空行を用意して行番号を返す（SUMIF を +1 拡張）。"""
    total = find_label_row(ws, config.TOTAL_LABEL)
    if total is None:
        raise ValueError(f"{ws.title}: 合計行が見つかりません")
    if total - 1 >= config.FIRST_DATA_ROW and _is_blank_row(ws, total - 1):
        return total - 1
    ret = total + 1
    # ① 結合解除（挿入後だと B 列に書けなくなる）
    for r in (total, ret):
        rng = f"A{r}:B{r}"
        if rng in {str(m) for m in ws.merged_cells.ranges}:
            ws.unmerge_cells(rng)
    # ② 挿入
    ws.insert_rows(total)
    new_row, total = total, total + 1
    # ③ 再結合 + SUMIF 範囲を書き直し
    _write_totals(ws, total)
    # ④ 縞模様は行番号の偶奇で塗る（隣行コピーはしない）。ルール範囲が足りなければ拡張
    style_data_row(ws, new_row)
    if total + 1 > config.RULE_LAST_ROW - 10:
        ensure_sheet_rules(ws, total + 1)
    return new_row


def add_entry(wb: Workbook, e: Entry) -> AddResult:
    """落札日（仕入日）の月シートに登録。同じ商品IDがあれば空欄だけ埋める。"""
    msgs: list[str] = []
    item_id = e.item_id
    if item_id:
        dup = find_by_item_id(wb, item_id)
        if dup:
            ws = wb[dup.sheet]
            filled = []
            if dup["送料"] in (None, "") and e.shipping is not None:
                ws.cell(dup.row, C["送料"], e.shipping)
                filled.append("送料")
            if not sanitize_invoice(dup["インボイス番号"]):
                t = resolve_invoice(wb, e.shop or dup["店舗名"] or "", e.invoice)
                if t:
                    ws.cell(dup.row, C["インボイス番号"], t)
                    filled.append("インボイス番号")
            msgs.append(f"商品ID {item_id} は登録済み（{dup.sheet} {dup.row}行目）")
            if filled:
                msgs.append("空欄を補完: " + "・".join(filled))
            return AddResult("filled" if filled else "duplicate", dup.sheet, dup.row, messages=msgs)

    invoice = resolve_invoice(wb, e.shop, e.invoice)
    if sanitize_invoice(e.invoice) != sanitize_invoice(invoice) and e.invoice:
        msgs.append(f"インボイス番号を {invoice or '空欄'} に補正（入力: {e.invoice}）")
    address = e.address
    if not address and e.shop_type == "実店舗":
        address = lookup_shop(wb, e.shop)[1]
    if e.shop_type == "ヤフオク個人":
        address = ""
    if not e.evidence and e.shop_type == "実店舗":
        e.evidence = config.MANUAL_EVIDENCE

    ws = wb[f"{e.date.month}月"]
    r = _append_row(ws)
    _write_entry(ws, r, e, invoice, address)
    style_data_row(ws, r)
    return AddResult("added", ws.title, r, invoice, msgs)


# ---------------------------------------------------------------- 集計

@dataclass
class MonthSummary:
    month: int
    count: int
    price: int
    shipping: int
    total: int
    returned_total: int


def month_summary(wb: Workbook, month: int) -> MonthSummary:
    """返品✓を除いた件数・金額・送料・合計（台帳の合計行と同じ定義）。"""
    cnt = price = ship = ret = 0
    for row in iter_rows(wb):
        if row.sheet != f"{month}月":
            continue
        p, s = to_int(row["金額"]), to_int(row["送料"])
        if row["返品"] == config.CHECK:
            ret += p + s
            continue
        cnt += 1
        price += p
        ship += s
    return MonthSummary(month, cnt, price, ship, price + ship, ret)
