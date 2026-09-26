from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from shiire import config, dashboard, freetext, mf_match
from shiire.files import LedgerLockedError, backup, check_locks, evidence_path
from shiire.ledger import Entry, add_entry, create_ledger, edit_ledger, find_label_row, month_summary

URL1 = "https://auctions.yahoo.co.jp/jp/auction/x1187654321"
URL2 = "https://page.auctions.yahoo.co.jp/jp/auction/1123456789"


@pytest.fixture
def ledger(tmp_path):
    path = tmp_path / config.LEDGER_NAME
    create_ledger(path)
    return path


def _yahoo(**kw):
    base = dict(date=date(2026, 7, 10), name="Nikon F3", price=30000, shipping=1200,
                shop="カメラ屋ストア", payment="クレジットカード AMEX 下4けた 0000", evidence=URL1)
    base.update(kw)
    return Entry(**base)


def test_new_ledger_structure(ledger):
    wb = load_workbook(ledger)
    assert wb.sheetnames == ["年間合計"] + [f"{m}月" for m in range(1, 13)]
    ws = wb["7月"]
    assert ws["A1"].value.startswith("2026年7月")
    assert [ws.cell(2, c).value for c in range(1, 16)] == config.HEADERS
    assert find_label_row(ws, config.TOTAL_LABEL) == 4
    assert ws["C4"].value == '=SUMIF(M3:M3,"<>✓",C3:C3)'
    assert ws["E5"].value == '=SUMIF(M3:M3,"✓",E3:E3)'
    assert {"A1:O1", "A4:B4", "A5:B5"} <= {str(m) for m in ws.merged_cells.ranges}
    assert ws.sheet_properties.tabColor.rgb == config.TAB_MONTH
    assert wb["年間合計"]["B9"].value == """=IFERROR(INDEX('7月'!C:C,MATCH("合　計",'7月'!A:A,0)),"")"""
    dv = ws.data_validations.dataValidation[0]
    assert str(dv.sqref) == "M3:O200" and dv.formula1 == '"✓,"'
    rules = {str(r.sqref): r.rules for r in ws.conditional_formatting}
    assert [r.priority for r in rules["A3:O200"]] == [1, 2]
    assert all(r.stopIfTrue for r in rules["A3:O200"])


def test_add_rows_extends_sumif_and_merges(ledger):
    with edit_ledger(ledger, do_backup=False) as wb:
        r1 = add_entry(wb, _yahoo())
        r2 = add_entry(wb, _yahoo(evidence=URL2, name="Canon AE-1", price=8000, shipping=None))
        r3 = add_entry(wb, _yahoo(evidence=URL2 + "0", name="Pentax", price=5000, shipping=800))
    assert (r1.row, r2.row, r3.row) == (3, 4, 5)
    ws = load_workbook(ledger)["7月"]
    assert find_label_row(ws, config.TOTAL_LABEL) == 6
    assert ws["C6"].value == '=SUMIF(M3:M5,"<>✓",C3:C5)'
    assert ws["D7"].value == '=SUMIF(M3:M5,"✓",D3:D5)'
    assert ws["E5"].value == '=IFERROR(C5+IF(D5="",0,D5),"")'
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert {"A6:B6", "A7:B7"} <= merged and "A4:B4" not in merged
    assert ws["A3"].value == datetime(2026, 7, 10) and ws["A3"].number_format == "yyyy/mm/dd"
    assert ws["A3"].fill.start_color.rgb == config.STRIPE_ODD
    assert ws["A4"].fill.start_color.rgb == config.STRIPE_EVEN
    assert ws["A5"].fill.start_color.rgb == config.STRIPE_ODD


def test_coupon_and_forbidden_invoice(ledger):
    with edit_ledger(ledger, do_backup=False) as wb:
        res = add_entry(wb, _yahoo(coupon=500, invoice="T4010401039979"))
    ws = load_workbook(ledger)["7月"]
    assert ws["C3"].value == 29500
    assert ws["L3"].value.startswith("クーポン使用 -500円")
    assert ws["J3"].value is None and res.invoice == ""


def test_duplicate_fills_blanks_only(ledger):
    with edit_ledger(ledger, do_backup=False) as wb:
        add_entry(wb, _yahoo(shipping=None))
        res = add_entry(wb, _yahoo(shipping=900, price=1, invoice="T1234567890123"))
    assert res.status == "filled"
    ws = load_workbook(ledger)["7月"]
    assert (ws["C3"].value, ws["D3"].value, ws["J3"].value) == (30000, 900, "T1234567890123")
    assert find_label_row(ws, config.TOTAL_LABEL) == 4


def test_invoice_reuse_and_fixed_mapping(ledger):
    with edit_ledger(ledger, do_backup=False) as wb:
        add_entry(wb, _yahoo(shop="カメラ屋ストア", invoice="T1111111111111"))
        r = add_entry(wb, _yahoo(shop="カメラ屋 ストア", evidence=URL2, date=date(2026, 8, 1)))
        k = add_entry(wb, Entry(date=date(2026, 8, 2), name="x", price=1, shop="カメラのキタムラ 新宿店",
                                shop_type="実店舗", invoice="T9999999999999"))
        b = add_entry(wb, Entry(date=date(2026, 8, 3), name="y", price=1, shop="バイセル 渋谷", shop_type="実店舗"))
    assert r.invoice == "T1111111111111"
    assert k.invoice == "T3490001010673"
    assert b.invoice == "T4010001074187"


def test_freetext_real_store(ledger):
    with edit_ledger(ledger, do_backup=False) as wb:
        add_entry(wb, Entry(date=date(2026, 6, 1), name="old", price=1, shop="キタムラ",
                            shop_type="実店舗", address="東京都新宿区1-1"))
        e = freetext.parse_line("キタムラ Nikon F100 仕入38500 送料1000 T3490001010673 7/11")
        res = add_entry(wb, e)
    assert (e.date, e.name, e.price, e.shipping) == (date(2026, 7, 11), "Nikon F100", 38500, 1000)
    ws = load_workbook(ledger)["7月"]
    row = res.row
    assert ws.cell(row, 9).value == "実店舗"
    assert ws.cell(row, 11).value == "手動入力（実店舗）"
    assert ws.cell(row, 7).value == "東京都新宿区1-1"


def test_summary_excludes_returns(ledger):
    with edit_ledger(ledger, do_backup=False) as wb:
        add_entry(wb, _yahoo())
        add_entry(wb, _yahoo(evidence=URL2, price=5000, shipping=500, returned=True))
    s = month_summary(load_workbook(ledger), 7)
    assert (s.count, s.price, s.shipping, s.total, s.returned_total) == (1, 30000, 1200, 31200, 5500)


def test_dashboard_regenerate_handles_string_dates(ledger, tmp_path):
    wb = load_workbook(ledger)
    add_entry(wb, _yahoo(name='back\\slash "q"'))
    add_entry(wb, _yahoo(evidence=URL2, date=date(2026, 1, 5)))
    wb["1月"]["A3"] = "2026/01/05"  # 1-2月は文字列日付
    html = tmp_path / config.DASHBOARD_NAME
    assert dashboard.regenerate(wb, html) == 2
    text = html.read_text(encoding="utf-8")
    assert '"date": "2026-01-05"' in text and 'back\\\\slash' in text
    assert dashboard.regenerate(wb, html) == 2  # 2回目は既存ファイルを置換
    assert text.count("const ENTRIES=") == 1


def test_mf_match(ledger, tmp_path):
    with edit_ledger(ledger, do_backup=False) as wb:
        add_entry(wb, _yahoo(invoice="T1111111111111"))                      # 31200 A
        add_entry(wb, _yahoo(evidence=URL2, shop="個人", price=2000, shipping=None, shop_type="ヤフオク個人"))
        add_entry(wb, _yahoo(evidence=URL2 + "1", shop="個人", price=3000, shipping=None, shop_type="ヤフオク個人"))
    csv = tmp_path / "mf.csv"
    csv.write_text("取引No,日付,金額\n2382,7/11,31200\n2457,7/12,5000\n2332,7/12,999\n", encoding="utf-8")
    ms = mf_match.match(load_workbook(ledger), mf_match.load_journals_csv(csv, 2026))
    assert [m.rank for m in ms] == ["A", "B", "C"]
    assert len(ms[1].rows) == 2
    assert mf_match.report(ms).endswith("合計3件=A1+B1+C1")
    assert mf_match.numbers_only(ms) == "2457.2332"


def test_locks_and_backup(ledger, tmp_path):
    for i in range(4):
        backup(ledger, now=datetime(2026, 9, 1, 0, 0, i))
    kept = sorted((tmp_path / "バックアップ").iterdir())
    assert [p.name for p in kept] == ["仕入台帳2026.backup_20260901_000002.xlsx",
                                      "仕入台帳2026.backup_20260901_000003.xlsx"]
    (tmp_path / ("~$" + ledger.name)).write_text("")
    with pytest.raises(LedgerLockedError):
        check_locks(ledger)


def test_evidence_path(tmp_path):
    p = evidence_path(tmp_path, "x1187654321", date(2026, 7, 10))
    assert p.relative_to(tmp_path).as_posix() == "仕入証憑　画像あり/2026年/7月/x1187654321_20260710.png"


def test_many_rows_extend_rules(ledger):
    wb = load_workbook(ledger)
    for i in range(210):
        add_entry(wb, _yahoo(evidence=f"{URL1}{i}"))
    ws = wb["7月"]
    assert find_label_row(ws, config.TOTAL_LABEL) == 213
    assert str(ws.data_validations.dataValidation[0].sqref) == "M3:O264"


def test_cards_file(tmp_path, monkeypatch):
    monkeypatch.delenv("SHIIRE_CARDS_FILE", raising=False)
    with pytest.raises(KeyError):
        config.resolve_payment("AMEX", tmp_path)
    (tmp_path / "cards.json").write_text('{"amex": "クレジットカード AMEX 下4けた 9999"}', encoding="utf-8")
    assert config.resolve_payment("amex", tmp_path) == "クレジットカード AMEX 下4けた 9999"
    assert config.resolve_payment("現金", tmp_path) == "現金"
    e = freetext.parse_line("キタムラ F100 仕入100 7/1 AMEX", base=tmp_path)
    assert e.payment.endswith("9999")


def test_screenshot_blank_check(tmp_path):
    from PIL import Image

    from shiire.screenshot_win import is_blank

    white = tmp_path / "w.png"
    Image.new("RGB", (50, 50), "white").save(white)
    text = tmp_path / "t.png"
    img = Image.new("RGB", (50, 50), "white")
    img.putpixel((5, 5), (0, 0, 0))
    img.save(text)
    assert is_blank(white) and not is_blank(text)


def test_freetext_separated_labels():
    e = freetext.parse_line("オールドレンズ愛好家 SMC PENTAX-M 50mm F1.4 仕入れ 8,280円 送料0円 1/4")
    assert (e.shop, e.name, e.price, e.shipping, e.date) == (
        "オールドレンズ愛好家", "SMC PENTAX-M 50mm F1.4", 8280, 0, date(2026, 1, 4))
    e = freetext.parse_line("キタムラ F100 仕入 38500 送料 1,000円 7/11")
    assert (e.name, e.price, e.shipping) == ("F100", 38500, 1000)
